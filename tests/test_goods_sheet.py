from decimal import Decimal

from openpyxl import Workbook

from app.domain.entities.imcp.result import VatInfo
from app.infrastructure.extraction.imcp.goods_sheet import GoodsSheet


_ANCHOR_AVG = (
    "Начальная (максимальная) цена договора (среднее арифметическое значение) "
    "составляет, руб. в т.ч. НДС 20%"
)
_ANCHOR_AVG_NO_VAT_RATE = (
    "Начальная (максимальная) цена договора (среднее арифметическое значение) "
    "составляет, руб."
)


def test_goods_sheet_is_me_by_title() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчет при закупке товаров"

    assert GoodsSheet.is_me(ws)


def test_goods_sheet_is_me_by_header_fallback() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист с расчетом"
    ws["B5"] = "Наименование товара"
    ws["H5"] = "НМЦ товара (среднее арифметическое значение), руб. в т.ч. НДС"

    assert GoodsSheet.is_me(ws)


def test_goods_sheet_extract_returns_sources_items_prices_and_imcp_avg() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист с расчетом"

    ws["A1"] = "№ п/п"
    ws["B1"] = "Наименование товара"
    ws["C1"] = "Цена единицы товара, руб."
    ws["E1"] = "НМЦ единицы товара, руб. с НДС"
    ws["F1"] = "Ед. изм"
    ws["G1"] = "Кол-во"
    ws["H1"] = "НМЦ товара (среднее арифметическое значение), руб. в т.ч. НДС"
    ws["I1"] = _ANCHOR_AVG

    ws["C2"] = "ООО Альфа"
    ws["D2"] = "ООО Бета"

    ws["A3"] = "1"
    ws["B3"] = "Фланг SoftSwitch"
    ws["C3"] = "2766,1548"
    ws["D3"] = "2974,36"
    ws["E3"] = "2904,958267"
    ws["F3"] = "шт"
    ws["G3"] = "4003"
    ws["H3"] = "11628547,94"

    ws["A4"] = "2"
    ws["B4"] = "Фланг Контакт"
    ws["C4"] = "12142,4892"
    ws["D4"] = "13056,44"
    ws["E4"] = "12751,78973"
    ws["F4"] = "шт"
    ws["G4"] = "160"
    ws["H4"] = "2040286,357"

    ws["B5"] = _ANCHOR_AVG
    ws["H5"] = "13668834,30"

    result = GoodsSheet(ws).extract()

    assert result is not None
    assert result.imcp_avg == Decimal("13668834.30")
    assert result.vat_rate_info == VatInfo(rate=20, included=True)
    assert result.goods_data is not None
    assert [organization.id for organization in result.goods_data.organizations] == [
        1,
        2,
    ]
    assert [organization.name for organization in result.goods_data.organizations] == [
        "ООО Альфа",
        "ООО Бета",
    ]
    assert [item.id for item in result.goods_data.goods_items] == [1, 2]
    assert [item.name for item in result.goods_data.goods_items] == [
        "Фланг SoftSwitch",
        "Фланг Контакт",
    ]
    assert [price.organization_id for price in result.goods_data.price_cells] == [
        1,
        2,
        1,
        2,
    ]
    assert [price.item_id for price in result.goods_data.price_cells] == [1, 1, 2, 2]
    assert [price.price for price in result.goods_data.price_cells] == [
        Decimal("2766.15"),
        Decimal("2974.36"),
        Decimal("12142.49"),
        Decimal("13056.44"),
    ]


def test_goods_sheet_extract_vat_rate_percent_none_when_not_in_label() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист с расчетом"

    ws["A1"] = "№ п/п"
    ws["B1"] = "Наименование товара"
    ws["C1"] = "Цена единицы товара, руб."
    ws["E1"] = "НМЦ единицы товара, руб. с НДС"
    ws["F1"] = "Ед. изм"
    ws["G1"] = "Кол-во"
    ws["H1"] = "НМЦ товара (среднее арифметическое значение), руб. в т.ч. НДС"
    ws["I1"] = _ANCHOR_AVG_NO_VAT_RATE

    ws["C2"] = "ООО Альфа"
    ws["D2"] = "ООО Бета"

    ws["A3"] = "1"
    ws["B3"] = "Товар А"
    ws["C3"] = "100"
    ws["D3"] = "100"
    ws["E3"] = "100"
    ws["F3"] = "шт"
    ws["G3"] = "1"
    ws["H3"] = "100"

    ws["B4"] = _ANCHOR_AVG_NO_VAT_RATE
    ws["H4"] = "100"

    result = GoodsSheet(ws).extract()

    assert result is not None
    assert result.imcp_avg == Decimal(100)
    assert result.vat_rate_info is None


def test_goods_sheet_extract_vat_when_imcp_avg_cell_empty() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист с расчетом"

    ws["A1"] = "№ п/п"
    ws["B1"] = "Наименование товара"
    ws["C1"] = "Цена единицы товара, руб."
    ws["E1"] = "НМЦ единицы товара, руб. с НДС"
    ws["F1"] = "Ед. изм"
    ws["G1"] = "Кол-во"
    ws["H1"] = "НМЦ товара (среднее арифметическое значение), руб. в т.ч. НДС"

    ws["C2"] = "ООО Альфа"

    ws["A3"] = "1"
    ws["B3"] = "Товар А"
    ws["C3"] = "100"
    ws["E3"] = "100"
    ws["F3"] = "шт"
    ws["G3"] = "1"
    ws["H3"] = "100"

    ws["B4"] = _ANCHOR_AVG

    result = GoodsSheet(ws).extract()

    assert result is not None
    assert result.imcp_avg is None
    assert result.vat_rate_info == VatInfo(rate=20, included=True)
