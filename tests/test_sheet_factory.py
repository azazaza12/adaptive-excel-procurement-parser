from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.infrastructure.extraction.excel_imcp.base_sheet import BaseSheet
from app.infrastructure.extraction.excel_imcp.sheet_factory import SheetFactory


class _WorksSheet(BaseSheet[str]):
    @classmethod
    def is_me(cls, ws: Worksheet) -> bool:
        return "работ" in ws.title.lower()

    def extract(self) -> str:
        return self.ws.title


class _GoodsSheet(BaseSheet[str]):
    @classmethod
    def is_me(cls, ws: Worksheet) -> bool:
        return "товар" in ws.title.lower()

    def extract(self) -> str:
        return self.ws.title


def test_sheet_factory_create_all_matches_sheet_classes() -> None:
    workbook = Workbook()
    workbook.active.title = "Общие данные"
    workbook.create_sheet("Расчет при закупке товаров")
    workbook.create_sheet("Расчет при закупке работ, услуг")

    factory = SheetFactory([_WorksSheet, _GoodsSheet])
    sheets = factory.create_all(workbook)

    assert [type(sheet) for sheet in sheets] == [_GoodsSheet, _WorksSheet]
    assert [sheet.ws.title for sheet in sheets] == [
        "Расчет при закупке товаров",
        "Расчет при закупке работ, услуг",
    ]


def test_sheet_factory_create_unique_keeps_first_match() -> None:
    workbook = Workbook()
    workbook.active.title = "Расчет при закупке товаров"
    workbook.create_sheet("Доп. товары")

    factory = SheetFactory([_GoodsSheet])
    unique_sheets = factory.create_unique(workbook)

    assert list(unique_sheets.keys()) == [_GoodsSheet]
    assert unique_sheets[_GoodsSheet].ws.title == "Расчет при закупке товаров"


def test_sheet_factory_create_unique_raises_on_duplicates() -> None:
    workbook = Workbook()
    workbook.active.title = "Расчет при закупке товаров"
    workbook.create_sheet("Доп. товары")

    factory = SheetFactory([_GoodsSheet], raise_on_duplicates=True)

    try:
        factory.create_unique(workbook)
        assert False
    except ValueError as exc:
        assert "Multiple sheets matched _GoodsSheet" == str(exc)
