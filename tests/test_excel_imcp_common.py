from decimal import Decimal

from app.infrastructure.extraction.imcp.common import (
    extract_amount_by_anchor,
    extract_amount_from_row,
    to_decimal,
)
from openpyxl import Workbook


def test_to_decimal_ignores_excel_error_values() -> None:
    assert to_decimal("#ДЕЛ/0!") is None
    assert to_decimal("#DIV/0!") is None
    assert to_decimal("#Н/Д") is None


def test_to_decimal_parses_regular_numeric_values() -> None:
    assert to_decimal("123,45") == Decimal("123.45")
    assert to_decimal(" 1 234,56 руб. ") == Decimal("1234.56")


def test_to_decimal_limits_precision_to_two_decimal_places() -> None:
    assert to_decimal("13572602.073333332") == Decimal("13572602.07")


def test_extract_amount_from_row_skips_anchor_cell_with_vat_percent() -> None:
    wb = Workbook()
    ws = wb.active
    ws.cell(
        row=1,
        column=1,
        value=(
            "Начальная (максимальная) цена договора "
            "(среднее арифметическое значение) составляет, "
            "руб. в т.ч. НДС 20%"
        ),
    )
    ws.cell(row=1, column=2, value="#DIV/0!")

    assert (
        extract_amount_from_row(
            ws,
            row_idx=1,
            anchor="начальн максимальн цен договор средн арифметическ",
        )
        is None
    )


def test_extract_amount_by_anchor_prefers_selected_column_value() -> None:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Наименование товара")
    ws.cell(row=1, column=5, value="НМЦ товара")
    ws.cell(
        row=3,
        column=1,
        value=(
            "Начальная (максимальная) цена договора "
            "(среднее арифметическое значение) составляет, руб. в т.ч. НДС 20%"
        ),
    )
    ws.cell(row=3, column=5, value="1 234,56")
    ws.cell(row=3, column=9, value="20")

    assert extract_amount_by_anchor(
        ws,
        "начальн максимальн цен договор средн арифметическ",
        preferred_col_idx=5,
    ) == Decimal("1234.56")


def test_extract_amount_by_anchor_falls_back_when_selected_column_invalid() -> None:
    wb = Workbook()
    ws = wb.active
    ws.cell(
        row=3,
        column=1,
        value=(
            "Начальная (максимальная) цена договора "
            "(среднее арифметическое значение) составляет, руб."
        ),
    )
    ws.cell(row=3, column=5, value="N/A")
    ws.cell(row=3, column=9, value="9 999,00")

    assert extract_amount_by_anchor(
        ws,
        "начальн максимальн цен договор средн арифметическ",
        preferred_col_idx=5,
    ) == Decimal("9999.00")
