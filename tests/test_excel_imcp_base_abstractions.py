from decimal import Decimal

from openpyxl import Workbook

from app.infrastructure.extraction.imcp.base_abstractions import (
    AnyOfAnchors,
    BaseList,
    BaseColumn,
    DecimalConverter,
    ExactAnchor,
    MatchAnchor,
    StringConverter,
)


def test_base_column_finds_column_and_converts_values() -> None:
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "Наименование поставщика"
    ws["C1"] = "Цена предложения"
    ws["B2"] = "  ООО Йота  "
    ws["C2"] = "1 000,50"

    name_column = BaseColumn(
        ws,
        MatchAnchor("наименование"),
        converter=StringConverter(),
        header_row_idx=1,
    )
    amount_column = BaseColumn(
        ws,
        MatchAnchor("цена"),
        converter=DecimalConverter(),
        header_row_idx=1,
    )

    assert name_column.col_idx == 2
    assert amount_column.col_idx == 3
    assert name_column.get(2) == "ООО Йота"
    assert amount_column.get(2) == Decimal("1000.50")


def test_base_column_returns_none_when_anchor_not_found() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Другой заголовок"
    ws["A2"] = "Значение"

    column = BaseColumn(ws, MatchAnchor("цена"), header_row_idx=1)

    assert column.col_idx is None
    assert column.get(2) is None


def test_base_list_builds_row_and_skips_fully_empty_rows() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Организация"
    ws["B1"] = "Сумма"
    ws["A2"] = "АО Тета"
    ws["B2"] = "15 500"
    ws["A3"] = None
    ws["B3"] = None

    org_column = BaseColumn(
        ws,
        ExactAnchor("Организация"),
        converter=StringConverter(),
        header_row_idx=1,
    )
    amount_column = BaseColumn(
        ws,
        ExactAnchor("Сумма"),
        converter=DecimalConverter(),
        header_row_idx=1,
    )

    rows = BaseList(
        org_column,
        amount_column,
        row_builder=lambda org, amount: {"org": org, "amount": amount},
    )

    assert rows.build_row(2) == {"org": "АО Тета", "amount": Decimal("15500")}
    assert rows.build_row(3) is None


def test_any_of_anchors_matches_when_one_anchor_matches() -> None:
    anchor = AnyOfAnchors(MatchAnchor("наименование"), MatchAnchor("цена"))

    assert anchor.matches("Цена предложения, руб.")
    assert not anchor.matches("Ед. изм.")
