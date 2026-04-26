from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from typing import cast

from openpyxl.worksheet.worksheet import Worksheet

from app.infrastructure.extraction.excel_imcp.base_table_extractor import (
    BaseTableExtractor,
)
from app.infrastructure.extraction.excel_imcp.models import (
    ExtractorConfig,
    OrganizationInfo,
    ParsingContext,
)


type ItemNameExtractor = Callable[[Worksheet, int, ParsingContext], str | None]
type MetricExtractor = Callable[[Worksheet, int, str, int, ParsingContext], object]
type OrgAmountExtractor = Callable[
    [Worksheet, int, OrganizationInfo, ParsingContext], object
]
type IsValidRowPredicate = Callable[["DynamicRowData"], bool]
type Aggregator[ResultT] = Callable[
    [ParsingContext, list["DynamicRowData"]], ResultT | None
]
type SetFieldHandler = Callable[["DynamicRowData", str, object], None]
type AddOrgAmountHandler = Callable[["DynamicRowData", OrganizationInfo, object], None]
type SetItemNameHandler = Callable[["DynamicRowData", str], None]
type EmptyRowFactory = Callable[[], "DynamicRowData"]


@dataclass(slots=True)
class DynamicRowData:
    name: str | None = None
    metrics: dict[str, object] = field(default_factory=dict)
    org_amounts: list[tuple[OrganizationInfo, object]] = field(default_factory=list)
    raw_data: dict[str, object] = field(default_factory=dict)


class DynamicTableExtractor[ResultT](BaseTableExtractor[DynamicRowData, ResultT]):
    def __init__(  # noqa: PLR0913
        self,
        config: ExtractorConfig,
        *,
        item_name_extractor: ItemNameExtractor | None = None,
        metric_extractor: MetricExtractor | None = None,
        org_amount_extractor: OrgAmountExtractor | None = None,
        is_valid_row: IsValidRowPredicate | None = None,
        aggregate: Aggregator[ResultT] | None = None,
        set_field: SetFieldHandler | None = None,
        add_org_amount: AddOrgAmountHandler | None = None,
        set_item_name: SetItemNameHandler | None = None,
        empty_row_factory: EmptyRowFactory | None = None,
    ) -> None:
        super().__init__(config)
        self._item_name_extractor = item_name_extractor
        self._metric_extractor = metric_extractor
        self._org_amount_extractor = org_amount_extractor
        self._is_valid_row_predicate = is_valid_row
        self._aggregator = aggregate
        self._set_field_handler = set_field
        self._add_org_amount_handler = add_org_amount
        self._set_item_name_handler = set_item_name
        self._empty_row_factory = empty_row_factory

    def extract_item_name(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        context: ParsingContext,
    ) -> str | None:
        if self._item_name_extractor is not None:
            return self._item_name_extractor(ws, row_idx, context)

        value = ws.cell(row=row_idx, column=context.header.item_col_idx).value
        if isinstance(value, str):
            stripped = value.strip()
            if stripped:
                return stripped
        return None

    def extract_metric(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        metric_name: str,
        col_idx: int,
        context: ParsingContext,
    ) -> object:
        if self._metric_extractor is not None:
            return self._metric_extractor(ws, row_idx, metric_name, col_idx, context)
        return ws.cell(row=row_idx, column=col_idx).value

    def extract_org_amount(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        organization: OrganizationInfo,
        context: ParsingContext,
    ) -> object:
        if self._org_amount_extractor is not None:
            return self._org_amount_extractor(ws, row_idx, organization, context)
        return ws.cell(row=row_idx, column=organization.col_idx).value

    def is_valid_row(self, row_data: DynamicRowData) -> bool:
        if self._is_valid_row_predicate is not None:
            return self._is_valid_row_predicate(row_data)
        return bool(row_data.metrics or row_data.org_amounts)

    def aggregate(
        self,
        context: ParsingContext,
        row_data_list: list[DynamicRowData],
    ) -> ResultT | None:
        if self._aggregator is not None:
            return self._aggregator(context, row_data_list)
        if not row_data_list:
            return None
        return cast("ResultT", row_data_list)

    def _set_field(
        self, row_data: DynamicRowData, field_name: str, value: object
    ) -> None:
        if self._set_field_handler is not None:
            self._set_field_handler(row_data, field_name, value)
            return
        row_data.metrics[field_name] = value
        row_data.raw_data[field_name] = value

    def _add_org_amount(
        self,
        row_data: DynamicRowData,
        organization: OrganizationInfo,
        amount: object,
    ) -> None:
        if self._add_org_amount_handler is not None:
            self._add_org_amount_handler(row_data, organization, amount)
            return
        row_data.org_amounts.append((organization, amount))

    def _set_item_name(self, row_data: DynamicRowData, item_name: str) -> None:
        if self._set_item_name_handler is not None:
            self._set_item_name_handler(row_data, item_name)
            return
        row_data.name = item_name
        row_data.raw_data["name"] = item_name

    def create_empty_row_data(self) -> DynamicRowData:
        if self._empty_row_factory is not None:
            return self._empty_row_factory()
        return DynamicRowData()
