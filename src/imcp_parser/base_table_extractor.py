from __future__ import annotations

import re
from abc import ABC, abstractmethod

from openpyxl.worksheet.worksheet import Worksheet

from app.infrastructure.extraction.imcp.common import (
    advance_empty_rows_series,
    find_anchor_columns_range_in_row_window,
    find_header_columns,
    is_service_subheader,
    is_stop_row_by_anchor_match,
    match_anchor,
)
from app.infrastructure.extraction.imcp.models import (
    ExtractorConfig,
    HeaderAnchorPattern,
    HeaderColumns,
    OrganizationInfo,
    ParsingContext,
)


class BaseTableExtractor[RowDataT, ResultT](ABC):
    def __init__(self, config: ExtractorConfig) -> None:
        self._config = config

    def extract(self, ws: Worksheet) -> ResultT | None:
        header = self._find_header(ws)
        if header is None:
            return None

        organizations = self._find_org_columns(ws, header=header)
        context = ParsingContext(header=header, organizations=organizations)
        row_data_list: list[RowDataT] = []

        for row_idx in range(self._data_start_row_idx(header=header), ws.max_row + 1):
            item_name = self.extract_item_name(ws, row_idx=row_idx, context=context)
            is_empty = item_name is None
            context.empty_rows_in_series, should_stop = advance_empty_rows_series(
                context.empty_rows_in_series,
                is_empty=is_empty,
                max_empty_rows_in_series=self._config.max_empty_rows_in_series,
            )
            if is_empty:
                if should_stop:
                    break
                continue

            if item_name is None:
                continue
            if self._is_stop_row(item_name):
                break

            row_data = self.create_empty_row_data()
            self._set_item_name(row_data, item_name)

            for metric_name, col_idx in context.header.metric_columns.items():
                metric = self.extract_metric(
                    ws,
                    row_idx=row_idx,
                    metric_name=metric_name,
                    col_idx=col_idx,
                    context=context,
                )
                self._set_field(row_data, metric_name, metric)

            for organization in context.organizations:
                amount = self.extract_org_amount(
                    ws,
                    row_idx=row_idx,
                    organization=organization,
                    context=context,
                )
                self._add_org_amount(row_data, organization, amount)

            if self.is_valid_row(row_data):
                row_data_list.append(row_data)

        return self.aggregate(context, row_data_list)

    def _find_header(self, ws: Worksheet) -> HeaderColumns | None:
        item_anchor = self._config.header_anchors.get("item_col_idx")
        if item_anchor is None:
            return None

        header_data = find_header_columns(ws, {"item_col_idx": item_anchor})
        if header_data is None:
            return None

        header_row_idx, columns = header_data
        item_col_idx = columns.get("item_col_idx")
        if item_col_idx is None:
            return None

        metric_anchors = {
            field_name: anchor
            for field_name, anchor in self._config.header_anchors.items()
            if field_name != "item_col_idx"
        }
        metric_columns = self._find_metric_columns(
            ws,
            header_row_idx=header_row_idx,
            item_col_idx=item_col_idx,
            metric_anchors=metric_anchors,
        )
        return HeaderColumns(
            header_row_idx=header_row_idx,
            item_col_idx=item_col_idx,
            metric_columns=metric_columns,
        )

    def _find_metric_columns(
        self,
        ws: Worksheet,
        *,
        header_row_idx: int,
        item_col_idx: int,
        metric_anchors: dict[str, HeaderAnchorPattern],
    ) -> dict[str, int]:
        metric_columns: dict[str, int] = {}
        if not metric_anchors:
            return metric_columns

        scan_end_row_idx = min(
            ws.max_row,
            header_row_idx + self._config.header_scan_depth,
        )
        for col_idx in range(item_col_idx + 1, ws.max_column + 1):
            for scan_row_idx in range(header_row_idx, scan_end_row_idx + 1):
                header_value = ws.cell(row=scan_row_idx, column=col_idx).value
                if not isinstance(header_value, str):
                    continue

                for field_name, anchor_pattern in metric_anchors.items():
                    if field_name in metric_columns:
                        continue
                    if col_idx in metric_columns.values():
                        continue
                    anchor_candidates = (
                        anchor_pattern
                        if isinstance(anchor_pattern, tuple)
                        else (anchor_pattern,)
                    )
                    if any(
                        match_anchor(header_value, anchor)
                        for anchor in anchor_candidates
                    ):
                        metric_columns[field_name] = col_idx
                        break

        return metric_columns

    def _find_org_columns(
        self,
        ws: Worksheet,
        *,
        header: HeaderColumns,
    ) -> list[OrganizationInfo]:
        if not self._config.org_header_anchors:
            return []

        metric_columns_set = set(header.metric_columns.values())
        scan_end_row_idx = min(
            ws.max_row,
            header.header_row_idx + self._config.header_scan_depth,
        )
        price_block_columns_range = self._find_price_block_columns_range(
            ws,
            header=header,
            header_row_idx=header.header_row_idx,
        )
        best_match: list[OrganizationInfo] = []

        for org_header_row_idx in range(header.header_row_idx, scan_end_row_idx + 1):
            if not self._is_org_header_row_candidate(
                ws=ws,
                header=header,
                row_idx=org_header_row_idx,
            ):
                continue
            current_match = self._find_org_columns_in_row(
                ws=ws,
                header=header,
                org_header_row_idx=org_header_row_idx,
                metric_columns_set=metric_columns_set,
                price_block_columns_range=price_block_columns_range,
            )
            if len(current_match) > len(best_match):
                best_match = current_match

        return best_match

    def _is_org_header_row_candidate(
        self,
        *,
        ws: Worksheet,
        header: HeaderColumns,
        row_idx: int,
    ) -> bool:
        if row_idx <= header.header_row_idx:
            return True

        # Organization headers may be on secondary header rows, but not on data rows.
        item_cell_value = ws.cell(row=row_idx, column=header.item_col_idx).value
        return not (isinstance(item_cell_value, str) and item_cell_value.strip())

    def _find_org_columns_in_row(
        self,
        *,
        ws: Worksheet,
        header: HeaderColumns,
        org_header_row_idx: int,
        metric_columns_set: set[int],
        price_block_columns_range: tuple[int, int] | None,
    ) -> list[OrganizationInfo]:
        organizations: list[OrganizationInfo] = []
        started = False

        for col_idx in range(header.item_col_idx + 1, ws.max_column + 1):
            raw_header_value = ws.cell(row=org_header_row_idx, column=col_idx).value
            should_continue, should_stop = self._resolve_org_column_skip(
                col_idx=col_idx,
                started=started,
                raw_header_value=raw_header_value,
                metric_columns_set=metric_columns_set,
                price_block_columns_range=price_block_columns_range,
            )
            if should_stop:
                break
            if should_continue:
                continue

            raw_value = ws.cell(row=org_header_row_idx, column=col_idx).value
            if not isinstance(raw_value, str):
                continue

            organizations.append(
                OrganizationInfo(
                    col_idx=col_idx,
                    name=raw_value.strip(),
                    id=len(organizations) + 1,
                )
            )
            started = True

        return organizations

    def _resolve_org_column_skip(
        self,
        *,
        col_idx: int,
        started: bool,
        raw_header_value: object,
        metric_columns_set: set[int],
        price_block_columns_range: tuple[int, int] | None,
    ) -> tuple[bool, bool]:
        if col_idx in metric_columns_set:
            return True, started

        in_price_block = self._is_in_price_block(
            col_idx=col_idx,
            price_block_columns_range=price_block_columns_range,
        )
        if price_block_columns_range is not None and not in_price_block:
            return True, started

        if not isinstance(raw_header_value, str):
            return True, started

        header_name = raw_header_value.strip()
        if self._is_skippable_org_header_name(
            header_name,
            require_org_anchor=price_block_columns_range is None,
        ):
            return True, started

        return False, False

    def _is_skippable_org_header_name(
        self,
        header_name: str,
        *,
        require_org_anchor: bool,
    ) -> bool:
        if not header_name:
            return True

        if any(
            match_anchor(header_name, anchor)
            for anchor in self._config.stop_column_anchors
        ):
            return True

        if is_service_subheader(
            header_name,
            self._config.price_block_service_subheaders,
        ):
            return True
        if _is_numeric_like_header(header_name):
            return True

        if not require_org_anchor:
            return False

        return not any(
            match_anchor(header_name, anchor)
            for anchor in self._config.org_header_anchors
        )

    def _find_price_block_columns_range(
        self,
        ws: Worksheet,
        *,
        header: HeaderColumns,
        header_row_idx: int,
    ) -> tuple[int, int] | None:
        if not self._config.price_block_parent_anchors:
            return None

        for anchor in self._config.price_block_parent_anchors:
            columns_range = find_anchor_columns_range_in_row_window(
                ws,
                anchor=anchor,
                start_row_idx=header_row_idx,
                scan_depth=self._config.header_scan_depth,
            )
            if columns_range is not None:
                min_col_idx, max_col_idx = columns_range
                if min_col_idx == max_col_idx:
                    expanded_range = self._expand_price_block_columns_range(
                        header=header,
                        anchor_col_idx=min_col_idx,
                    )
                    if expanded_range is not None:
                        return expanded_range
                return columns_range
        return None

    def _expand_price_block_columns_range(
        self,
        *,
        header: HeaderColumns,
        anchor_col_idx: int,
    ) -> tuple[int, int] | None:
        right_bound = min(
            (
                col_idx
                for col_idx in header.metric_columns.values()
                if col_idx > header.item_col_idx
            ),
            default=None,
        )
        if right_bound is None:
            return None
        max_col_idx = right_bound - 1
        if max_col_idx < header.item_col_idx + 1:
            return None
        min_col_idx = header.item_col_idx + 1
        if anchor_col_idx < min_col_idx or anchor_col_idx > max_col_idx:
            return None
        return min_col_idx, max_col_idx

    def _is_in_price_block(
        self,
        *,
        col_idx: int,
        price_block_columns_range: tuple[int, int] | None,
    ) -> bool:
        if price_block_columns_range is None:
            return False
        min_col_idx, max_col_idx = price_block_columns_range
        return min_col_idx <= col_idx <= max_col_idx

    def _is_stop_row(self, item_name: str) -> bool:
        if not self._config.stop_row_anchors:
            return False
        return is_stop_row_by_anchor_match(item_name, self._config.stop_row_anchors)

    def _data_start_row_idx(self, *, header: HeaderColumns) -> int:
        return header.header_row_idx + 1

    @abstractmethod
    def extract_item_name(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        context: ParsingContext,
    ) -> str | None:
        raise NotImplementedError

    @abstractmethod
    def extract_metric(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        metric_name: str,
        col_idx: int,
        context: ParsingContext,
    ) -> object:
        raise NotImplementedError

    @abstractmethod
    def extract_org_amount(
        self,
        ws: Worksheet,
        *,
        row_idx: int,
        organization: OrganizationInfo,
        context: ParsingContext,
    ) -> object:
        raise NotImplementedError

    @abstractmethod
    def is_valid_row(self, row_data: RowDataT) -> bool:
        raise NotImplementedError

    @abstractmethod
    def aggregate(
        self,
        context: ParsingContext,
        row_data_list: list[RowDataT],
    ) -> ResultT | None:
        raise NotImplementedError

    @abstractmethod
    def _set_field(self, row_data: RowDataT, field_name: str, value: object) -> None:
        raise NotImplementedError

    @abstractmethod
    def _add_org_amount(
        self,
        row_data: RowDataT,
        organization: OrganizationInfo,
        amount: object,
    ) -> None:
        raise NotImplementedError

    @abstractmethod
    def _set_item_name(self, row_data: RowDataT, item_name: str) -> None:
        raise NotImplementedError

    @abstractmethod
    def create_empty_row_data(self) -> RowDataT:
        raise NotImplementedError


_NUMERIC_HEADER_RE = re.compile(
    r"^\s*[-+]?\d[\d\s\xa0]*(?:[.,]\d+)?\s*$",
    flags=re.UNICODE,
)


def _is_numeric_like_header(value: str) -> bool:
    return bool(_NUMERIC_HEADER_RE.fullmatch(value))
