from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities.excel.goods_result import (
    GoodsItemResult,
    GoodsOrgAmountResult,
    GoodsOrgResult,
    GoodsOrgsAmountsResult,
)
from app.infrastructure.extraction.excel_imcp.cell_converters import (
    DecimalCellConverter,
    StringCellConverter,
)
from app.infrastructure.extraction.excel_imcp.common import (
    extract_amount_by_anchor,
    find_metric_column_idx,
)
from app.infrastructure.extraction.excel_imcp.dynamic_table_extractor import (
    DynamicRowData,
)
from app.infrastructure.extraction.excel_imcp.extractor_builder import ExtractorBuilder
from app.infrastructure.extraction.excel_imcp.models import (
    GoodsSheetParseResult,
    OrganizationInfo,
    ParsingContext,
)


_GOODS_TABLE_ITEM_HEADER_ANCHOR = "наименовани товар"
_GOODS_TABLE_STOP_ROW_ANCHORS = (
    "инструкц",
    "начальн максимальн цен договор",
)
_GOODS_TABLE_STOP_COLUMN_ANCHORS = (
    "ед изм",
    "кол",
)
_GOODS_PRICE_BLOCK_PARENT_ANCHORS = ("цена единицы товара",)
_GOODS_PRICE_BLOCK_SERVICE_SUBHEADERS = (
    "руб",
    "за ед",
    "общая",
    "ед изм",
    "кол",
)
_GOODS_TABLE_ORG_HEADER_ANCHORS = (
    "орг",
    "ооо",
    "ао",
    "ип",
    "зао",
    "пао",
)
_GOODS_HEADER_SCAN_DEPTH = 2
_METRIC_UNIT_IMCP_ANCHORS = ("нмц единиц",)
_METRIC_UNIT_OF_MEASURE_ANCHORS = ("ед изм",)
_METRIC_QUANTITY_ANCHORS = ("кол", "всего")
_METRIC_ITEM_IMCP_ANCHORS = ("нмц товар", "нмц")
_MAX_EMPTY_ROWS_IN_SERIES = 3
_IMCP_ANCHOR_AVG = "начальн максимальн цен договор средн арифметическ"
_DECIMAL_CONVERTER = DecimalCellConverter()
_STRING_CONVERTER = StringCellConverter()


def _extract_item_name(
    ws: Worksheet, row_idx: int, context: ParsingContext
) -> str | None:
    value = ws.cell(row=row_idx, column=context.header.item_col_idx).value
    return _STRING_CONVERTER.convert(value)


def _extract_metric(
    ws: Worksheet,
    row_idx: int,
    metric_name: str,
    col_idx: int,
    context: ParsingContext,
) -> object:
    _ = context
    value = ws.cell(row=row_idx, column=col_idx).value
    if metric_name == "unit_of_measure":
        return _STRING_CONVERTER.convert(value)
    return _DECIMAL_CONVERTER.convert(value)


def _extract_org_amount(
    ws: Worksheet,
    row_idx: int,
    organization: OrganizationInfo,
    context: ParsingContext,
) -> object:
    _ = context
    col_idx = organization.col_idx
    value = ws.cell(row=row_idx, column=col_idx).value
    return _DECIMAL_CONVERTER.convert(value)


def _is_valid_goods_row(row_data: DynamicRowData) -> bool:
    has_metrics = any(value is not None for value in row_data.metrics.values())
    has_org_amounts = any(amount is not None for _, amount in row_data.org_amounts)
    return has_metrics or has_org_amounts


def _aggregate_goods(
    context: ParsingContext,
    row_data_list: list[DynamicRowData],
) -> GoodsOrgsAmountsResult | None:
    if not context.organizations or not row_data_list:
        return None

    goods_items: list[GoodsItemResult] = []
    cells: list[GoodsOrgAmountResult] = []
    for row_data in row_data_list:
        if row_data.name is None:
            continue
        item_id = len(goods_items) + 1
        goods_items.append(
            GoodsItemResult(
                id=item_id,
                name=row_data.name,
                unit_of_measure=row_data.metrics.get("unit_of_measure"),
                unit_imcp=row_data.metrics.get("unit_imcp"),
                quantity=row_data.metrics.get("quantity"),
                item_imcp=row_data.metrics.get("item_imcp"),
            )
        )
        for organization, amount in row_data.org_amounts:
            if amount is None:
                continue
            cells.append(
                GoodsOrgAmountResult(
                    organization_id=organization.id,
                    item_id=item_id,
                    amount=amount,
                )
            )

    if not goods_items:
        return None

    return GoodsOrgsAmountsResult(
        organizations=[
            GoodsOrgResult(
                id=organization.id,
                name=organization.name,
            )
            for organization in context.organizations
        ],
        goods_items=goods_items,
        cells=cells,
    )


def extract_goods_sheet(
    ws: Worksheet,
) -> GoodsSheetParseResult:
    item_imcp_col_idx = find_metric_column_idx(
        ws,
        item_anchor=_GOODS_TABLE_ITEM_HEADER_ANCHOR,
        metric_anchor=_METRIC_ITEM_IMCP_ANCHORS,
        header_scan_depth=_GOODS_HEADER_SCAN_DEPTH,
    )
    extractor = (
        ExtractorBuilder[GoodsOrgsAmountsResult]()
        .with_header_anchors(
            {
                "item_col_idx": _GOODS_TABLE_ITEM_HEADER_ANCHOR,
                "unit_imcp": _METRIC_UNIT_IMCP_ANCHORS[0],
                "unit_of_measure": _METRIC_UNIT_OF_MEASURE_ANCHORS[0],
                "quantity": _METRIC_QUANTITY_ANCHORS[0],
                "item_imcp": _METRIC_ITEM_IMCP_ANCHORS,
            }
        )
        .with_stop_row_anchors(*_GOODS_TABLE_STOP_ROW_ANCHORS)
        .with_stop_column_anchors(*_GOODS_TABLE_STOP_COLUMN_ANCHORS)
        .with_header_scan_depth(_GOODS_HEADER_SCAN_DEPTH)
        .with_price_block_parent_anchors(*_GOODS_PRICE_BLOCK_PARENT_ANCHORS)
        .with_price_block_service_subheaders(*_GOODS_PRICE_BLOCK_SERVICE_SUBHEADERS)
        .with_org_header_anchors(*_GOODS_TABLE_ORG_HEADER_ANCHORS)
        .with_max_empty_rows_in_series(_MAX_EMPTY_ROWS_IN_SERIES)
        .with_item_name_extractor(_extract_item_name)
        .with_metric_extractor(_extract_metric)
        .with_org_amount_extractor(_extract_org_amount)
        .with_is_valid_row(_is_valid_goods_row)
        .with_aggregator(_aggregate_goods)
        .build()
    )
    return GoodsSheetParseResult(
        imcp_avg=extract_amount_by_anchor(
            ws,
            _IMCP_ANCHOR_AVG,
            preferred_col_idx=item_imcp_col_idx,
        ),
        orgs_amounts=extractor.extract(ws),
    )
