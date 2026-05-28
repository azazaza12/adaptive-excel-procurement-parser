from abc import ABC, abstractmethod
from collections.abc import Sequence
from dataclasses import dataclass
from typing import ClassVar

from openpyxl.worksheet.worksheet import Worksheet

from app.infrastructure.extraction.imcp.base_abstractions import BaseAnchor


@dataclass(slots=True)
class HeaderCell:
    row_idx: int
    col_idx: int


class BaseSheet[SheetResultT](ABC):
    _registry: ClassVar[list[type[BaseSheet]]] = []

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws

    def __init_subclass__(cls, **kwargs: object) -> None:
        super().__init_subclass__(**kwargs)
        BaseSheet._registry.append(cls)

    @classmethod
    def registry(cls) -> tuple[type[BaseSheet], ...]:
        return tuple(BaseSheet._registry)

    @classmethod
    @abstractmethod
    def is_me(cls, ws: Worksheet) -> bool:
        pass

    @abstractmethod
    def extract(self) -> SheetResultT | None:
        pass

    def row_has_anchor(self, *, row_idx: int, anchor: BaseAnchor) -> bool:
        for col_idx in range(1, self.ws.max_column + 1):
            value = self.ws.cell(row=row_idx, column=col_idx).value
            if anchor.matches(value):
                return True
        return False

    def _find_cell_by(
        self, anchor: BaseAnchor, *, row_idx: int | None = None
    ) -> HeaderCell | None:
        if row_idx is not None:
            for col_idx in range(1, self.ws.max_column + 1):
                value = self.ws.cell(row=row_idx, column=col_idx).value
                if anchor.matches(value):
                    return HeaderCell(row_idx=row_idx, col_idx=col_idx)
            return None
        for row_idx_iter, row in enumerate(self.ws.iter_rows(), start=1):
            for cell in row:
                value = cell.value
                if anchor.matches(value):
                    return HeaderCell(row_idx=row_idx_iter, col_idx=cell.column)
        return None

    def _find_header_by(self, anchor: BaseAnchor) -> HeaderCell | None:
        return self._find_cell_by(anchor)

    def _has_title_anchors(self, anchor: BaseAnchor) -> bool:
        return anchor.matches(self.ws.title)

    def _has_table_header(self, anchors: Sequence[BaseAnchor]) -> bool:
        headers = [self._find_header_by(anchor) for anchor in anchors]
        return bool(all(header is not None for header in headers))
