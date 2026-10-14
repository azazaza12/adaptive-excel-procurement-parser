from dataclasses import dataclass
from decimal import Decimal

from openpyxl.worksheet.worksheet import Worksheet

from .base_abstractions import AnyOfAnchors, MatchAnchor
from .base_sheet import BaseSheet
from .common import (
    parse_number,
    parse_vat_info_from_cell_text,
    to_decimal,
    to_stripped_str,
)
from .entities import VatInfo
from .entities.works_result import WorksOrgAmountResult
from .models import WorksSheetParseResult


_WORKS_SHEET_TITLE_ANCHORS = AnyOfAnchors(MatchAnchor("работ"), MatchAnchor("услуг"))
_MAX_EMPTY_ROWS_IN_SERIES = 3

_ORG_HEADER_ANCHOR = MatchAnchor("организац")
_INDEX_HEADER_ANCHORS = AnyOfAnchors(MatchAnchor("№ п п"), MatchAnchor("номер"))
_AMOUNT_HEADER_WORKS_OR_SERVICES_ANCHORS = AnyOfAnchors(
    MatchAnchor("стоимост работ"), MatchAnchor("стоимост услуг")
)
_IMCP_ROW_ANCHOR = MatchAnchor("начальн максимальн цен договор")
_IMCP_AVG_ROW_ANCHOR = MatchAnchor("начальн максимальн цен договор средн арифметическ")
_IMCP_LIMIT_ROW_ANCHOR = MatchAnchor("лимит финансирован")
_INSTRUCTION_ROW_ANCHOR = MatchAnchor("инструкц")
_STOP_ROW_ANCHORS = AnyOfAnchors(_IMCP_ROW_ANCHOR, _INSTRUCTION_ROW_ANCHOR)


@dataclass(slots=True)
class WorksHeader:
    row_idx: int
    index_col_idx: int | None
    org_col_idx: int
    amount_col_idx: int


class WorksSheet(BaseSheet[WorksSheetParseResult]):
    @classmethod
    def is_me(cls, ws: Worksheet) -> bool:
        sheet = cls(ws)
        if sheet._has_title_anchors(_WORKS_SHEET_TITLE_ANCHORS):
            return True
        return sheet._has_table_header(
            (_ORG_HEADER_ANCHOR, _AMOUNT_HEADER_WORKS_OR_SERVICES_ANCHORS)
        )

    def extract(self) -> WorksSheetParseResult | None:
        header = None
        index_header = self._find_header_by(_INDEX_HEADER_ANCHORS)
        org_header = self._find_header_by(_ORG_HEADER_ANCHOR)
        amount_header = self._find_header_by(_AMOUNT_HEADER_WORKS_OR_SERVICES_ANCHORS)
        if (
            org_header
            and amount_header
            and index_header
            and org_header.row_idx == amount_header.row_idx == index_header.row_idx
        ):
            header = WorksHeader(
                row_idx=org_header.row_idx,
                index_col_idx=index_header.col_idx,
                org_col_idx=org_header.col_idx,
                amount_col_idx=amount_header.col_idx,
            )
        if header is None:
            return None
        amount_header_vat_info = parse_vat_info_from_cell_text(
            self.ws.cell(row=header.row_idx, column=header.amount_col_idx).value
        )

        org_amount_rows, org_block_end_row_idx = (
            self.extract_org_amount_rows_until_imcp(header)
        )
        imcp_avg, imcp_limit, imcp_avg_row_vat_info = self.extract_imcp_from_imcp_block(
            imcp_start_row_idx=org_block_end_row_idx + 1
            if org_block_end_row_idx is not None
            else None,
            amount_col_idx=header.amount_col_idx,
        )
        return WorksSheetParseResult(
            imcp_avg=imcp_avg,
            imcp_limit=imcp_limit,
            orgs_amounts=org_amount_rows,
            amount_header_vat_info=amount_header_vat_info,
            imcp_avg_row_vat_info=imcp_avg_row_vat_info,
        )

    def extract_org_amount_rows_until_imcp(
        self, header: WorksHeader
    ) -> tuple[list[WorksOrgAmountResult] | None, int | None]:
        organizations: list[WorksOrgAmountResult] = []
        org_block_end_row_idx: int | None = None

        for row_idx in range(header.row_idx + 1, self.ws.max_row + 1):
            if header.index_col_idx is None:
                break
            raw_item_number = self.ws.cell(
                row=row_idx,
                column=header.index_col_idx,
            ).value
            parsed_item_number = parse_number(raw_item_number)
            if parsed_item_number is None:
                break

            name = to_stripped_str(
                self.ws.cell(row=row_idx, column=header.org_col_idx).value
            )
            amount = to_decimal(
                self.ws.cell(row=row_idx, column=header.amount_col_idx).value
            )
            if name is None or amount is None:
                continue

            organizations.append(
                WorksOrgAmountResult(
                    id=len(organizations) + 1,
                    name=name,
                    amount=amount,
                )
            )
            org_block_end_row_idx = row_idx
        return organizations or None, org_block_end_row_idx

    def extract_imcp_from_imcp_block(
        self,
        *,
        imcp_start_row_idx: int | None,
        amount_col_idx: int,
    ) -> tuple[Decimal | None, Decimal | None, VatInfo | None]:
        imcp_avg: Decimal | None = None
        imcp_limit: Decimal | None = None
        vat_rate_info: VatInfo | None = None
        if imcp_start_row_idx is None:
            return imcp_avg, imcp_limit, vat_rate_info

        for row_idx in range(imcp_start_row_idx, self.ws.max_row + 1):
            label_cell = self._find_cell_by(_IMCP_AVG_ROW_ANCHOR, row_idx=row_idx)
            if label_cell is not None:
                imcp_avg = to_decimal(
                    self.ws.cell(row=row_idx, column=amount_col_idx).value,
                )
                label_value = self.ws.cell(
                    row=label_cell.row_idx, column=label_cell.col_idx
                ).value
                vat_rate_info = parse_vat_info_from_cell_text(label_value)

            if imcp_limit is None and self.row_has_anchor(
                row_idx=row_idx, anchor=_IMCP_LIMIT_ROW_ANCHOR
            ):
                imcp_limit = to_decimal(
                    self.ws.cell(row=row_idx, column=amount_col_idx).value
                )

            if (
                imcp_avg is not None
                and imcp_limit is not None
                and vat_rate_info is not None
            ):
                break
            if self.row_has_anchor(row_idx=row_idx, anchor=_INSTRUCTION_ROW_ANCHOR):
                break

        return imcp_avg, imcp_limit, vat_rate_info
