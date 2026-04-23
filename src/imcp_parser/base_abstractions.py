from abc import ABC, abstractmethod
from decimal import Decimal
from typing import Any, Callable, Generic, TypeVar

from openpyxl.worksheet.worksheet import Worksheet

from app.infrastructure.extraction.excel_imcp.common import match_anchor, norm_text, to_decimal

ConvertedT = TypeVar("ConvertedT")
RowT = TypeVar("RowT")


class BaseAnchor(ABC):
    @abstractmethod
    def matches(self, value: str) -> bool:
        pass


class MatchAnchor(BaseAnchor):
    def __init__(self, anchor: str) -> None:
        self._anchor = anchor

    def matches(self, value: str) -> bool:
        if not isinstance(value, str):
            return False
        return match_anchor(value, self._anchor)


class ExactAnchor(BaseAnchor):
    def __init__(self, anchor: str) -> None:
        self._anchor_norm = norm_text(anchor)

    def matches(self, value: str) -> bool:
        if not isinstance(value, str):
            return False
        return norm_text(value) == self._anchor_norm


class AnyOfAnchors(BaseAnchor):
    def __init__(self, *anchors: BaseAnchor) -> None:
        self._anchors = anchors

    def matches(self, value: str) -> bool:
        return any(anchor.matches(value) for anchor in self._anchors)


class BaseConverter(ABC, Generic[ConvertedT]):
    @abstractmethod
    def convert(self, value: Any) -> ConvertedT | None:
        pass


class IdentityConverter(BaseConverter[Any]):
    def convert(self, value: Any) -> Any:
        return value


class StringConverter(BaseConverter[str]):
    def convert(self, value: Any) -> str | None:
        if not isinstance(value, str):
            return None
        stripped = value.strip()
        return stripped or None


class DecimalConverter(BaseConverter[Decimal]):
    def convert(self, value: Any) -> Decimal | None:
        return to_decimal(value)


class BaseColumn(Generic[ConvertedT]):
    def __init__(
        self,
        ws: Worksheet,
        anchor: BaseAnchor,
        *,
        converter: BaseConverter[ConvertedT] | None = None,
        header_row_idx: int | None = None,
    ) -> None:
        self._ws = ws
        self._anchor = anchor
        self._converter = converter or IdentityConverter()
        self.col_idx = self._find_col_idx(header_row_idx=header_row_idx)

    def _find_col_idx(self, *, header_row_idx: int | None) -> int | None:
        if header_row_idx is not None:
            rows = self._ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx)
        else:
            rows = self._ws.iter_rows()

        for row in rows:
            for cell in row:
                if self._anchor.matches(cell.value):
                    return cell.column
        return None

    def get(self, row_idx: int) -> ConvertedT | None:
        if self.col_idx is None:
            return None
        cell_value = self._ws.cell(row=row_idx, column=self.col_idx).value
        return self._converter.convert(cell_value)


class BaseList(Generic[RowT]):
    def __init__(
        self,
        *columns: BaseColumn[Any],
        row_builder: Callable[..., RowT],
    ) -> None:
        self._columns = columns
        self._row_builder = row_builder

    def build_row(self, row_idx: int) -> RowT | None:
        values = [column.get(row_idx) for column in self._columns]
        if all(value is None for value in values):
            return None
        return self._row_builder(*values)
