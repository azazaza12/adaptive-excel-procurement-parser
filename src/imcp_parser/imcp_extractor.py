from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities.imcp.result import IMCPResult
from app.domain.exceptions import (
    IMCPNotFoundError,
)
from app.domain.ports.excel_imcp_extractor import IExcelIMCPExtractor
from app.infrastructure.extraction.imcp.base_sheet import BaseSheet
from app.infrastructure.extraction.imcp.common import norm_text
from app.infrastructure.extraction.imcp.result_updater import update_result


_ORGANIZATIONS_HEADS = (
    "организация",
    "контрагент",
    "поставщик",
    "подрядчик",
    "исполнитель",
)

MIN_ORGANIZATION_NAME_LENGTH = 2


def _normalize_name(name: str) -> str:
    return " ".join(name.strip().lower().replace("ё", "е").split())


def _is_placeholder(name: str) -> bool:
    norm_name = norm_text(name) or ""
    if not norm_name or len(norm_name) <= MIN_ORGANIZATION_NAME_LENGTH:
        return True

    tokens = norm_name.split()
    head = tokens[0]
    tail = tokens[1:]

    if head in _ORGANIZATIONS_HEADS:
        return len(tail) == 0 or all(t.isdigit() for t in tail)
    if head.startswith("организац"):
        return len(tail) == 0 or all(t.isdigit() for t in tokens)
    if head == "орг" and tail:
        marker = tail[0]
        rest = tail[1:]
        if marker in {"ция", "я"}:
            return len(rest) == 0 or all(t.isdigit() for t in rest)
    return False


def _name_score(name: str) -> tuple[int, int, int]:
    n = _normalize_name(name)
    not_placeholder = 0 if _is_placeholder(n) else 1
    words = len(n.split())
    length = len(n)
    return (not_placeholder, words, length)


def _pick_normalized_name(names: list[str]) -> str:
    non_empty = [n for n in names if n and n.strip()]
    if not non_empty:
        return ""
    return max(non_empty, key=_name_score)


def _normalize_organization_names(result: IMCPResult) -> IMCPResult:
    if (
        result.goods_data
        and result.goods_data.organizations
        and result.works_orgs_amounts
    ):
        names_by_id: dict[int, list[str]] = {}
        normalized_names_by_id: dict[int, str] = {}
        for org in result.goods_data.organizations:
            names_by_id.setdefault(org.id, []).append(org.name)
        for org in result.works_orgs_amounts:
            names_by_id.setdefault(org.id, []).append(org.name)

        for org_id, names in names_by_id.items():
            normalized_name = _pick_normalized_name(names)
            normalized_names_by_id[org_id] = normalized_name
        for i, org in enumerate(result.goods_data.organizations):
            if org.id in normalized_names_by_id:
                result.goods_data.organizations[i].name = normalized_names_by_id[org.id]
        for i, org in enumerate(result.works_orgs_amounts):
            if org.id in normalized_names_by_id:
                result.works_orgs_amounts[i].name = normalized_names_by_id[org.id]
    return result


def _build_sheet(ws: Worksheet) -> BaseSheet | None:
    for sheet_cls in BaseSheet.registered():
        if sheet_cls.is_me(ws):
            return sheet_cls(ws)
    return None


class OpenpyxlExcelIMCPExtractor(IExcelIMCPExtractor):
    def extract_imcp(self, *, content: bytes) -> IMCPResult:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=False)

        try:
            result = IMCPResult()
            has_supported_data = False
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet = _build_sheet(ws)
                if sheet is None:
                    continue
                parsed = sheet.extract()
                if parsed is not None:
                    update_result(parsed, result)
                    has_supported_data = True
            if not has_supported_data:
                raise IMCPNotFoundError
            _normalize_organization_names(result)
            return result
        finally:
            wb.close()
