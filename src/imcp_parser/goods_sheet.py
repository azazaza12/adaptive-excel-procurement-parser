from dataclasses import dataclass
from decimal import Decimal

from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities.imcp.goods_result import (
    GoodsData,
    GoodsItem,
    PriceCell,
    Source,
)
from app.infrastructure.extraction.imcp.base_abstractions import (
    AllOfAnchors,
    AnyOfAnchors,
    MatchAnchor,
    NotContainsAnchor,
)
from app.infrastructure.extraction.imcp.base_sheet import BaseSheet
from app.infrastructure.extraction.imcp.common import (
    parse_number,
    to_decimal,
    to_stripped_str,
)
from app.infrastructure.extraction.imcp.models import GoodsSheetParseResult


_INDEX_HEADER_ANCHORS = AnyOfAnchors(MatchAnchor("№ п п"), MatchAnchor("номер"))
_GOODS_SHEET_TITLE_ANCHOR = MatchAnchor("товар")
_ITEM_NAME_HEADER_ANCHOR = MatchAnchor("наименова товар")
_PRICE_BLOCK_PARENT_ANCHORS = MatchAnchor("цена единицы товара")
_ITEM_IMCP_ANCHORS = AllOfAnchors(
    AnyOfAnchors(MatchAnchor("нмц товар"), MatchAnchor("нмц средн арифм")),
    NotContainsAnchor("един"),
)
_UNIT_IMCP_ANCHORS = MatchAnchor("нмц единиц")
_UNIT_OF_MEASURE_ANCHORS = MatchAnchor("ед изм")
_QUANTITY_ANCHORS = AnyOfAnchors(MatchAnchor("кол"), MatchAnchor("всего"))
_IMCP_ANCHOR_AVG = MatchAnchor("начальн максимальн цен договор средн арифметическ")


@dataclass(slots=True)
class ParsedSource:
    col_idx: int
    name: str


@dataclass(slots=True)
class GoodsHeader:
    start_row_idx: int
    end_row_idx: int
    index_col_idx: int
    item_name_col_idx: int
    price_block_start_col_idx: int
    unit_imcp_col_idx: int
    unit_of_measure_col_idx: int | None
    quantity_col_idx: int
    item_imcp_col_idx: int


class GoodsSheet(BaseSheet):
    @classmethod
    def is_me(cls, ws: Worksheet) -> bool:
        sheet = cls(ws)
        if sheet._has_title_anchors(_GOODS_SHEET_TITLE_ANCHOR):
            return True
        return sheet._has_table_header((_ITEM_IMCP_ANCHORS, _ITEM_NAME_HEADER_ANCHOR))

    def extract(self) -> GoodsSheetParseResult | None:
        header = None
        index_header = self._find_header_by(_INDEX_HEADER_ANCHORS)
        item_name_header = self._find_header_by(_ITEM_NAME_HEADER_ANCHOR)
        price_block_header = self._find_header_by(_PRICE_BLOCK_PARENT_ANCHORS)
        unit_imcp_header = self._find_header_by(_UNIT_IMCP_ANCHORS)
        unit_of_measure_header = self._find_header_by(_UNIT_OF_MEASURE_ANCHORS)
        quantity_header = self._find_header_by(_QUANTITY_ANCHORS)
        item_imcp_header = self._find_header_by(_ITEM_IMCP_ANCHORS)

        if (
            index_header
            and item_name_header
            and price_block_header
            and unit_imcp_header
            and quantity_header
            and item_imcp_header
        ):
            header = GoodsHeader(
                start_row_idx=index_header.row_idx,
                end_row_idx=price_block_header.row_idx + 1,
                index_col_idx=index_header.col_idx,
                item_name_col_idx=item_name_header.col_idx,
                price_block_start_col_idx=price_block_header.col_idx,
                unit_imcp_col_idx=unit_imcp_header.col_idx,
                unit_of_measure_col_idx=unit_of_measure_header.col_idx
                if unit_of_measure_header
                else None,
                quantity_col_idx=quantity_header.col_idx,
                item_imcp_col_idx=item_imcp_header.col_idx,
            )

        if header is None:
            return None

        price_block_end_col_idx = unit_imcp_header.col_idx - 1
        parsed_sources = self._extract_sources(
            header_row_idx=price_block_header.row_idx,
            start_col_idx=price_block_header.col_idx,
            end_col_idx=price_block_end_col_idx,
        )

        sources, col_idx_to_source_id = self._map_sources(parsed_sources)
        goods_items, prices, block_end_row_idx = self.extract_goods_items_until_imcp(
            header, parsed_sources, col_idx_to_source_id
        )
        if prices is None or goods_items is None:
            return None
        imcp_avg = self.extract_imcp_from_imcp_block(
            block_end_row_idx=block_end_row_idx,
            item_imcp_col_idx=header.item_imcp_col_idx,
        )
        return GoodsSheetParseResult(
            goods_data=GoodsData(
                sources=sources, goods_items=goods_items, price_cells=prices
            ),
            imcp_avg=imcp_avg,
        )

    def _extract_sources(
        self,
        *,
        header_row_idx: int,
        start_col_idx: int,
        end_col_idx: int,
    ) -> list[ParsedSource]:
        sources: list[ParsedSource] = []

        for col_idx in range(start_col_idx, end_col_idx + 1):
            raw_value = self.ws.cell(row=header_row_idx + 1, column=col_idx).value
            name = raw_value.strip() if isinstance(raw_value, str) else ""
            if not name:
                continue
            sources.append(ParsedSource(col_idx=col_idx, name=name))

        return sources

    def _map_sources(
        self, parsed_sources: list[ParsedSource]
    ) -> tuple[list[Source], dict[int, int]]:
        sources: list[Source] = []
        col_idx_to_source_id: dict[int, int] = {}
        for i, parsed in enumerate(parsed_sources, start=1):
            source_id = i
            sources.append(Source(id=source_id, name=parsed.name))
            col_idx_to_source_id[parsed.col_idx] = source_id
        return sources, col_idx_to_source_id

    def extract_prices_for_item(
        self,
        *,
        row_idx: int,
        item_id: int,
        parsed_sources: list[ParsedSource],
        col_idx_to_source_id: dict[int, int],
    ) -> list[PriceCell]:
        prices: list[PriceCell] = []
        for parsed_source in parsed_sources:
            price = to_decimal(
                self.ws.cell(row=row_idx, column=parsed_source.col_idx).value
            )
            if price is None:
                continue
            prices.append(
                PriceCell(
                    source_id=col_idx_to_source_id[parsed_source.col_idx],
                    item_id=item_id,
                    price=price,
                )
            )
        return prices

    def extract_goods_items_until_imcp(
        self,
        header: GoodsHeader,
        parsed_sources: list[ParsedSource],
        col_idx_to_source_id: dict[int, int],
    ) -> tuple[list[GoodsItem] | None, list[PriceCell], int | None]:
        goods_items: list[GoodsItem] = []
        prices: list[PriceCell] = []
        block_end_row_idx: int | None = None

        for row_idx in range(header.end_row_idx + 1, self.ws.max_row + 1):
            if header.index_col_idx is None:
                break
            raw_iter_number = self.ws.cell(
                row=row_idx, column=header.index_col_idx
            ).value
            parsed_item_number = parse_number(raw_iter_number)
            if parsed_item_number is None:
                break
            name = to_stripped_str(
                self.ws.cell(row=row_idx, column=header.item_name_col_idx).value
            )
            unit_imcp = to_decimal(
                self.ws.cell(row=row_idx, column=header.unit_imcp_col_idx).value
            )
            unit_of_measure = to_stripped_str(
                self.ws.cell(row=row_idx, column=header.unit_of_measure_col_idx).value
            )
            quantity = to_decimal(
                self.ws.cell(row=row_idx, column=header.quantity_col_idx).value
            )
            item_imcp = to_decimal(
                self.ws.cell(row=row_idx, column=header.item_imcp_col_idx).value
            )
            if (
                name is None
                or unit_imcp is None
                or quantity is None
                or item_imcp is None
            ):
                continue
            goods_items.append(
                GoodsItem(
                    id=parsed_item_number,
                    name=name,
                    unit_of_measure=unit_of_measure,
                    unit_imcp=unit_imcp,
                    quantity=quantity,
                    item_imcp=item_imcp,
                )
            )
            prices_for_item = self.extract_prices_for_item(
                row_idx=row_idx,
                item_id=parsed_item_number,
                parsed_sources=parsed_sources,
                col_idx_to_source_id=col_idx_to_source_id,
            )
            prices.extend(prices_for_item)
            block_end_row_idx = row_idx
        goods_items = goods_items or None
        return goods_items, prices, block_end_row_idx

    def extract_imcp_from_imcp_block(
        self, *, block_end_row_idx: int | None, item_imcp_col_idx: int
    ) -> Decimal | None:
        imcp_avg = None
        if block_end_row_idx is None:
            return imcp_avg
        for row_idx in range(block_end_row_idx + 1, self.ws.max_row + 1):
            if self.row_has_anchor(row_idx=row_idx, anchor=_IMCP_ANCHOR_AVG):
                imcp_avg = to_decimal(
                    self.ws.cell(row=row_idx, column=item_imcp_col_idx).value
                )
            if imcp_avg is not None:
                break
        return imcp_avg
