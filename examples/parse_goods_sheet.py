"""Minimal end-to-end example: build a goods sheet in memory and parse it.

Run from the repository root after `pip install -e .`:

    python examples/parse_goods_sheet.py
"""

from io import BytesIO

from openpyxl import Workbook

from imcp_parser import OpenpyxlExcelIMCPExtractor


def build_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover - openpyxl always creates a sheet
        raise RuntimeError("no active worksheet")

    ws.title = "Товары"
    ws["A1"] = "№ п/п"
    ws["B1"] = "Наименование товара"
    ws["C1"] = "Цена единицы товара"
    ws["C2"] = "ООО Альфа"
    ws["D2"] = "ООО Бета"
    ws["E1"] = "НМЦ единицы товара"
    ws["F1"] = "Количество"
    ws["G1"] = "НМЦ товара"
    ws["A3"] = 1
    ws["B3"] = "Бумага А4"
    ws["C3"] = "100,50"
    ws["D3"] = "101,50"
    ws["E3"] = "101,00"
    ws["F3"] = 10
    ws["G3"] = "1010,00"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    result = OpenpyxlExcelIMCPExtractor().extract_imcp(content=build_workbook())
    goods = result.goods_data

    print("Организации:")
    for org in goods.organizations:
        print(f"    {org.id} {org.name}")

    print("Позиции:")
    for item in goods.goods_items:
        print(f"    {item.id} {item.name}")
        print(f"        кол-во: {item.quantity}, НМЦ: {item.item_imcp}")

    print("Цены:")
    for cell in goods.price_cells:
        print(f"    орг. {cell.organization_id} / поз. {cell.item_id} — {cell.price}")


if __name__ == "__main__":
    main()
